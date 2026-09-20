#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
cad_model_code_compare.py - 用本机广联达工程量 Excel 做梁对象级自检。

用途：
  不把“梁平法图自动中心线粗算”直接与广联达总额比较，而是先按编号集合判断
  CAD 图纸最可能对应模型中的哪一层，再输出每个梁编号的自动方量缺口。

    cad_model_code_compare.py \
        --model-xlsx 1#.xlsx \
        --auto-json 3_5f_beam_auto.json -o 梁构件模型对照

结果仍用于研发自检，不作为结算或翻样口径。
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

SKILL_DIR = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = SKILL_DIR / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))


HEAD_RE = re.compile(
    r"(WKL|XL|LL|JZL|JGL|TXL|KL|DL|TL|WL|ZHL|QL|L)\s*\d{0,3}[A-Za-z]?",
    re.I)


def norm_beam_code(value: Any) -> str:
    """模型与 CAD 编号统一成小写无括号/无分隔符。"""
    text = str(value or "").replace(" ", "")
    m = HEAD_RE.search(text)
    if not m:
        return re.sub(r"\W+", "", text).lower()
    code = m.group(0).lower()
    return code


def load_model_beam_rows(path: Path, sheet_name: str | None = None) -> dict[str, dict[str, float]]:
    """读取广联达“绘图输入工程量汇总表-梁”，保留每层每编号方量。"""
    try:
        import openpyxl
    except Exception as exc:
        raise SystemExit(f"需要 openpyxl 读取模型 Excel：{exc}") from exc
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet_name is None:
        sheet_name = next(
            (n for n in wb.sheetnames
             if "绘图输入工程量汇总表" in n and n.endswith("-梁")),
            None)
    if not sheet_name or sheet_name not in wb.sheetnames:
        raise SystemExit(f"模型工作簿中未找到梁表：{sheet_name or wb.sheetnames}")
    ws = wb[sheet_name]
    out: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or row[0] is None:
            continue
        floor = str(row[0]).strip()
        name = str(row[1] or "").strip()
        if not name or name in ("小计", "合计"):
            continue
        grade = str(row[2] or "").strip().upper()
        cat = str(row[3] or "").strip()
        if grade != "C30" or cat != "梁":
            continue
        try:
            volume = float(row[4] or 0)
        except (TypeError, ValueError):
            continue
        if volume > 0:
            out[floor][norm_beam_code(name)] += volume
    return {k: dict(v) for k, v in out.items()}


def load_auto(rows: list[dict[str, Any]],
              quantity_field: str = "volume_m3") -> dict[str, float]:
    """保留 CAD 已识别编号集合，即使无中心线也要参与图纸匹配。"""
    out: dict[str, float] = defaultdict(float)
    for r in rows or []:
        out[norm_beam_code(r.get("base_code") or r.get("code") or "")] += \
            float(r.get(quantity_field) or 0)
    return dict(out)


def load_auto_topology(data: dict[str, Any]) -> dict[str, float]:
    """读取梁拓扑台账的编号汇总，保留未归属段供审计。"""
    registry = data.get("run_registry") or {}
    code_rows = registry.get("code_quantities") or {}
    out: dict[str, float] = {}
    for key, row in code_rows.items():
        code = "__unassigned__" if key == "__unassigned__" else norm_beam_code(
            row.get("code") or key)
        out[code] = float(row.get("estimated_volume_m3") or 0.0)
    return out


def extract_auto_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    """兼容 cad_beam_objects 平铺 rows 和 cad_structure_model 的 members.beams。"""
    rows = data.get("rows")
    if isinstance(rows, list):
        return rows
    members = data.get("members") or {}
    beams = members.get("beams")
    return beams if isinstance(beams, list) else []


def nested_value(data: dict[str, Any], path: str | None) -> float | None:
    if not path:
        return None
    value: Any = data
    for part in path.split("."):
        if not isinstance(value, dict):
            return None
        value = value.get(part)
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def code_score(auto: set[str], model: set[str]) -> tuple[float, int, int, int]:
    shared = len(auto & model)
    if not auto or not model:
        return 0.0, shared, len(auto), len(model)
    score = 0.6 * shared / len(auto) + 0.4 * shared / len(model)
    return score, shared, len(auto), len(model)


def render_md(items: list[dict[str, Any]], quantity_field: str) -> str:
    o = ["# CAD 梁编号与广联达模型构件级对照", ""]
    o.append("本表用于研发自检：先判断 CAD 图纸应匹配哪一模型层，再按编号看自动梁方量缺口。")
    o.append(f"CAD 数量字段：`{quantity_field}`。")
    if any(item.get("auto_topology_m3") is not None for item in items):
        o.append("另有统一拓扑台账方量，作为编号级粗算之外的独立口径。")
    o.append("")
    has_topology = any(
        item.get("auto_topology_m3") is not None for item in items)
    o.append("| CAD 图纸 | 匹配模型层组 | 匹配分 | 模型梁m3 | CAD自动m3 | "
             + ("CAD拓扑m3 | 拓扑差异% | " if has_topology else "")
             + "差异% | CAD独有 | 模型独有 |")
    o.append("|---|---|---:|---:|---:|"
             + ("---:|---:|" if has_topology else "")
             + "---:|---|---|")
    for it in items:
        model_m3 = it["model_m3"]
        auto_m3 = it["auto_m3"]
        diff = (auto_m3 - model_m3) / model_m3 * 100 if model_m3 else None
        topology = it.get("auto_topology_m3")
        topology_diff = (
            (topology - model_m3) / model_m3 * 100
            if topology is not None and model_m3 else None
        )
        topology_cells = (
            f" {topology:.3f} | "
            f"{'-' if topology_diff is None else round(topology_diff, 1)}% |"
            if has_topology else ""
        )
        o.append(
            f"| {it['floor_label']} | {it['best_floor_text']} | "
            f"{it['best_score']:.3f} | {model_m3:.3f} | {auto_m3:.3f} |"
            f"{topology_cells} "
            f"{'-' if diff is None else round(diff, 1)}% | "
            f"{'、'.join(it['extra_codes'][:8]) or '-'} | "
            f"{'、'.join(it['missing_codes'][:8]) or '-'} |")
        o.append("")
        if it["code_rows"]:
            o.append("#### " + it["floor_label"])
            o.append("")
            o.append("| 编号 | 模型m3 | CAD编号级m3 | CAD拓扑m3 | "
                     "拓扑差异m3 | CAD对象 |")
            o.append("|---|---:|---:|---:|---:|---:|")
            for code_row in it["code_rows"]:
                diff_v = code_row["model"] - code_row["auto"]
                o.append(f"| {code_row['code']} | {code_row['model']:.3f} | "
                         f"{code_row['auto']:.3f} | "
                         f"{code_row['topology']:.3f} | "
                         f"{code_row['model']-code_row['topology']:+.3f} | "
                         f"{code_row['objects']} |")
            o.append("")
    o += ["", "## 说明", "",
          "- 匹配模型层组按 CAD 图纸已识别梁编号集合自动匹配；同编号的多层会合并显示。",
          "- 自动 JSON 中 `TL/LL` 来自楼梯/连梁详图时，不会出现在主梁平法图，所以会被列为模型独有。",
          "- 当 CAD 图纸同时包含多个结构层（如某层梁图里混入 WKL），编号集合会自动暴露映射问题。",
          ""]
    return "\n".join(o)


def csv_lines(items: list[dict[str, Any]], quantity_field: str) -> list[list[str]]:
    headers = ["CAD图纸", "匹配模型层组", "匹配分", "模型层", "模型梁m3", "CAD自动m3",
               "CAD拓扑m3", "拓扑差异%", "编号级差异m3", "编号级拓扑差异m3",
               "CAD数量字段",
               "CAD独有编号", "模型独有编号", "候选模型层"]
    out = [headers]
    for it in items:
        model_m3 = it["model_m3"]
        diff = (it["auto_m3"] - model_m3) / model_m3 * 100 if model_m3 else ""
        topology = it.get("auto_topology_m3")
        topology_diff = (
            (topology - model_m3) / model_m3 * 100
            if topology is not None and model_m3 else ""
        )
        out.append([
            it["floor_label"], it["best_floor_text"], f"{it['best_score']:.3f}",
            it["model_floor_source"], f"{model_m3:.3f}", f"{it['auto_m3']:.3f}",
            "" if topology is None else f"{topology:.3f}",
            f"{topology_diff:.1f}" if topology_diff != "" else "",
            "",
            "",
            quantity_field,
            "、".join(it["extra_codes"]), "、".join(it["missing_codes"]),
            it["candidate_text"],
        ])
        for code_row in it["code_rows"]:
            out.append([f"  {code_row['code']}", "", "", it["best_floor_text"],
                        f"{code_row['model']:.3f}", f"{code_row['auto']:.3f}",
                        f"{code_row['topology']:.3f}", "",
                        f"{code_row['model']-code_row['auto']:+.3f}",
                        f"{code_row['model']-code_row['topology']:+.3f}",
                        "", "", "", ""])
    return out


def _floor_label_span(floors: list[str]) -> str:
    """把同编号组写成可读范围，尽量保留原始层名。"""
    if len(floors) <= 3:
        return "、".join(floors)
    first = floors[0]
    last = floors[-1]
    def seq_no(f: str) -> int | None:
        m = re.search(r"第\s*(\d+)\s*层", f)
        return int(m.group(1)) if m else None
    a = seq_no(first)
    b = seq_no(last)
    if a is not None and b is not None and b - a == len(floors) - 1:
        return f"{first}~{last}"
    return f"{first} 等 {len(floors)} 层"


def _floor_order_key(floor: str) -> tuple[int, str]:
    m = re.search(r"第\s*(\d+)\s*层", floor)
    if m:
        return (0, int(m.group(1)), floor)
    if floor.startswith("基础层"):
        return (-1, 0, floor)
    if floor.startswith("首层"):
        return (0, 1, floor)
    return (1, 0, floor)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--model-xlsx", required=True, help="广联达导出的 xlsx")
    ap.add_argument("--model-sheet", default=None, help="可选：梁表 sheet 名")
    ap.add_argument("--auto-json", nargs="+", required=True)
    ap.add_argument("--quantity-field", default="volume_m3",
                    choices=["volume_m3", "gross_volume_m3", "clear_volume_m3"],
                    help="自动 JSON 中参与对比的梁方量字段")
    ap.add_argument(
        "--auto-total-field",
        default=None,
        help="额外读取自动 JSON 的点路径总结方量，如 "
             "run_registry.quantities.estimated_volume_m3",
    )
    ap.add_argument("--model-floor", action="append", default=[],
                    help="可选覆盖：图纸名=模型层，可重复")
    ap.add_argument("--floor-range", default=None,
                    help="可选：指定模型层范围，如 第2层-第5层")
    ap.add_argument("-o", "--out", default="cad_model_code_compare")
    args = ap.parse_args()

    model = load_model_beam_rows(Path(args.model_xlsx), args.model_sheet)
    overrides: dict[str, str] = {}
    for item in args.model_floor:
        if "=" in item:
            k, v = item.split("=", 1)
            overrides[k.strip()] = v.strip()

    items = []
    for auto_path in args.auto_json:
        data = json.loads(Path(auto_path).read_text(encoding="utf-8"))
        label = str(
            data.get("floor_label")
            or (data.get("floor") or {}).get("label")
            or ""
        )
        rows = extract_auto_rows(data)
        auto = load_auto(rows, args.quantity_field)
        topology_by_code = load_auto_topology(data)
        auto_set = set(auto)
        candidates = []
        for floor, model_rows in model.items():
            score, _shared, _ac, _mc = code_score(auto_set, set(model_rows))
            if score > 0:
                candidates.append((score, floor))
        candidates.sort(key=lambda x: (-x[0], x[1]))
        if candidates:
            best_score = candidates[0][0]
            best_floors = sorted(
                (f for s, f in candidates if abs(s - best_score) < 1e-9),
                key=_floor_order_key)
        else:
            best_score = 0.0
            best_floors = []
        best_floor_text = _floor_label_span(best_floors) if best_floors else "-"
        best_floor = overrides.get(label) or (best_floors[0] if best_floors else "-")
        if best_floor not in model:
            model_m3 = 0.0
            missing = sorted(auto_set)
            extra = []
            model_rows = {}
        else:
            model_rows = model[best_floor]
            model_m3 = sum(model_rows.values())
            missing = sorted(auto_set - set(model_rows))
            extra = sorted(set(model_rows) - auto_set)
        merged_codes = sorted(
            set(auto) | set(topology_by_code)
            | set(model_rows if best_floor in model else {}))
        code_rows = []
        code_objects = Counter()
        for r in rows:
            code_objects[norm_beam_code(r.get("base_code") or r.get("code") or "")] += 1
        for code in merged_codes:
            code_rows.append({
                "code": code,
                "model": round(float(model_rows.get(code, 0.0) or 0), 3),
                "auto": round(auto.get(code, 0.0), 3),
                "topology": round(
                    float(topology_by_code.get(code, 0.0) or 0), 3),
                "objects": code_objects.get(code, 0),
            })
        items.append({
            "floor_label": label,
            "best_floor_text": best_floor_text,
            "best_model_floor": best_floor,
            "best_score": round(best_score, 3),
            "candidate_text": "; ".join(f"{f} {s:.2f}" for s, f in candidates[:4]),
            "model_floor_source": "override" if label in overrides else "auto",
            "model_m3": round(model_m3, 3),
            "auto_m3": round(sum(auto.values()), 3),
            "auto_topology_m3": (
                round(nested_value(data, args.auto_total_field), 3)
                if nested_value(data, args.auto_total_field) is not None
                else None
            ),
            "extra_codes": missing,
            "missing_codes": extra,
            "code_rows": code_rows,
        })

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.with_suffix(".md").write_text(
        render_md(items, args.quantity_field), encoding="utf-8")
    with out.with_suffix(".csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(csv_lines(items, args.quantity_field))
    out.with_suffix(".json").write_text(
        json.dumps(items, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"已写出 {out.with_suffix('.md')} / {out.with_suffix('.csv')} / "
          f"{out.with_suffix('.json')}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
